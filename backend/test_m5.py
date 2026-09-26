import requests
import json
import os
import docx
import openpyxl
from fpdf import FPDF

BASE_URL = "http://localhost:8000/api"

print("--- REGRESSION TESTS ---")
print("1. Health:", requests.get(f"{BASE_URL}/health").json())
print("2. System Status:", requests.get(f"{BASE_URL}/system/status").json())
print("3. Models:", requests.get(f"{BASE_URL}/models").json())

print("\n--- M5 TESTS ---")

# Helpers to create documents
def create_test_files():
    # 1. Normal PDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, "Hello from normal PDF", ln=1)
    pdf.add_page()
    pdf.cell(200, 10, "Page 2 text", ln=1)
    pdf.output("test_normal.pdf")
    
    # 2. No-text PDF (simulated by creating a PDF with no text elements or drawing)
    pdf_empty = FPDF()
    pdf_empty.add_page()
    pdf_empty.output("test_empty.pdf")
    
    # 3. TXT
    with open("test.txt", "w", encoding="utf-8") as f:
        f.write("Hello from TXT")
        
    # 4. DOCX
    doc = docx.Document()
    doc.add_paragraph("Paragraph one in DOCX")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0,0).text = "A"
    table.cell(0,1).text = "B"
    table.cell(1,0).text = "C"
    table.cell(1,1).text = "D"
    doc.save("test.docx")
    
    # 5. CSV
    with open("test.csv", "w", encoding="utf-8") as f:
        f.write("H1,H2\n1,2\n3,4")
        
    # 6. XLSX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Col1", "Col2"])
    ws.append([1, 2])
    wb.save("test.xlsx")

def upload_file(path, content_type):
    with open(path, "rb") as f:
        return requests.post(f"{BASE_URL}/documents/upload", files={"file": (path, f, content_type)}).json()

create_test_files()
print("Test files created.")

docs = []
docs.append(("PDF Normal", upload_file("test_normal.pdf", "application/pdf")))
docs.append(("PDF Empty", upload_file("test_empty.pdf", "application/pdf")))
docs.append(("TXT", upload_file("test.txt", "text/plain")))
docs.append(("DOCX", upload_file("test.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")))
docs.append(("CSV", upload_file("test.csv", "text/csv")))
docs.append(("XLSX", upload_file("test.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")))

print("\n--- CONTENT EXTRACTION ---")
for name, meta in docs:
    doc_id = meta.get("document_id")
    if doc_id:
        resp = requests.get(f"{BASE_URL}/documents/{doc_id}/content")
        print(f"[{name}] -> Status: {resp.status_code}")
        # print first 150 chars of text
        res_json = resp.json()
        print(f"  Requires OCR: {res_json.get('requires_ocr')}")
        print(f"  Status: {res_json.get('content_status')}")
        text = res_json.get("content", {}).get("text", "")
        print(f"  Extracted text (preview): {text[:150]}")

print("\n--- ERROR HANDLING ---")
print("Nonexistent doc content:", requests.get(f"{BASE_URL}/documents/does-not-exist/content").json())

print("\n--- CLEANUP ---")
for _, meta in docs:
    if meta.get("document_id"):
        requests.delete(f"{BASE_URL}/documents/{meta['document_id']}")
        
os.remove("test_normal.pdf")
os.remove("test_empty.pdf")
os.remove("test.txt")
os.remove("test.docx")
os.remove("test.csv")
os.remove("test.xlsx")

print("Done.")
