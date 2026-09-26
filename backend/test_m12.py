
import os
from fastapi.testclient import TestClient
from app.main import app
from app.core.config import settings
from docx import Document

client = TestClient(app)

MOCK_DOC_DATA = {
    "document_type": "permit_extension_authorization",
    "organization": "Sovereign Industrial Operations",
    "title": "PERMIT-TO-WORK (PTW) EXTENSION AUTHORIZATION",
    "document_id": "DOC-12345",
    "extension_id": "EXT-2074614A",
    "permit_reference": "PTW-003",
    "equipment": "HX-301",
    "classification": "CONFIDENTIAL",
    "timestamp": "2026-09-25T10:00:00Z",
    "status": "DRAFT",
    "human_review_required": True,
    "permit_details": {
        "work_area": "Area 51",
        "requested_extension": "4 hours",
        "current_status": "EXPIRED",
        "requested_by": "John Doe",
        "approver_role": "Shift-In-Charge"
    },
    "agent_checks": [
        {
            "check_id": "check_permit_conflicts",
            "name": "PERMIT CONFLICT CHECK",
            "status": "PASS",
            "summary": "No conflicts found.",
            "evidence": "No conflicting permits were identified for the equipment.",
            "reference": "OISD-116",
            "is_blocking": False
        },
        {
            "check_id": "check_isolation_status",
            "name": "ISOLATION STATUS CHECK",
            "status": "CONFIRMED",
            "summary": "Isolation is valid.",
            "evidence": "Required isolation status was confirmed.",
            "reference": "SOP-123",
            "is_blocking": False
        },
        {
            "check_id": "check_gas_test_validity",
            "name": "GAS TEST VALIDITY CHECK",
            "status": "EXPIRED",
            "summary": "Gas test validity expired.",
            "evidence": "The latest gas test exceeded the permitted validity window.",
            "reference": "REG-456",
            "is_blocking": True
        }
    ],
    "findings": [
        {
            "finding_id": "Finding 1",
            "observation": "The latest gas test is outside the permitted validity window.",
            "evidence": "Gas test was done 24 hours ago.",
            "reference": "Internal SOP"
        }
    ],
    "regulatory_references": [
        {
            "document": "Internal SOP",
            "section": "3.4.2",
            "page": "15",
            "requirement": "Gas test must be fresh",
            "application": "Ensures no toxic gas build up."
        }
    ],
    "blocking_conditions": [
        {
            "condition": "Gas test validity expired.",
            "status": "BLOCKING",
            "reason": "Required validity window has been exceeded.",
            "evidence": "Test is 24 hours old.",
            "required_action": "A fresh gas test must be conducted before extension authorization can be considered.",
            "reference": "REG-456"
        }
    ],
    "required_actions": [
        {
            "description": "Conduct a fresh gas test.",
            "responsible_role": "Gas Tester"
        }
    ],
    "approval": {
        "approver_role": "Shift-In-Charge",
        "decision": "PENDING",
        "requires_signature": True
    },
    "provenance": {
        "agent": "Tuffy",
        "task_id": "task_m12_test",
        "source_documents": ["doc_1", "doc_2"],
        "knowledge_sources": []
    }
}

print("--- M12 TESTS ---")

def validate_docx(file_path):
    assert os.path.exists(file_path), "File does not exist"
    assert os.path.getsize(file_path) > 0, "File is empty"
    
    doc = Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs])
    
    # 2. DOCUMENT STRUCTURE (Required sections)
    assert "AI AGENT SAFETY VERIFICATION & DRAFT SUMMARY" in text
    assert "PERMIT DETAILS" in text
    assert "AGENT CHECK RESULTS SUMMARY" in text
    assert "FINDINGS SUMMARY" in text
    assert "REGULATORY REFERENCE SUMMARY" in text
    assert "BLOCKING CONDITIONS" in text
    assert "ACTION REQUIRED" in text
    assert "AUTHORITY AND APPROVAL" in text
    assert "DOCUMENT PROVENANCE" in text
    
    # 8. Human Review warning
    assert "AI-DRAFTED — Requires human review and sign-off" in text
    assert "DOCUMENT IS NOT VALID WITHOUT REQUIRED HUMAN REVIEW AND SIGN-OFF" in text

def test_generate_and_validate_document():
    from app.services.document_generator import document_generator_service
    task_id = "test_task_123"
    
    file_path = document_generator_service.generate_approval_document(task_id, MOCK_DOC_DATA)
    
    # 1. BASIC INDUSTRIAL DOCUMENT
    validate_docx(file_path)
    
    doc = Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs])
    
    # 3. AGENT CHECKS
    assert "PASS" in text
    assert "CONFIRMED" in text
    assert "EXPIRED" in text
    
    # 4. BLOCKING CONDITIONS
    assert "BLOCKING CONDITION #1" in text
    assert "Gas test validity expired." in text
    assert "NO EXTENSION SHALL BE GRANTED UNTIL ALL BLOCKING CONDITIONS ARE RESOLVED." in text
    
    # 6. REGULATORY REFERENCES
    assert "Internal SOP" in text
    
    # 8. HUMAN APPROVAL (Default state)
    assert "NOT YET APPROVED" in text
    
    print("TEST 1-4, 6, 8, 12 PASSED: Basic Document Generation and Validation")

def test_missing_data():
    from app.services.document_generator import document_generator_service
    task_id = "test_task_456"
    
    empty_data = {
        "organization": "Test Org"
    }
    
    file_path = document_generator_service.generate_approval_document(task_id, empty_data)
    validate_docx(file_path)
    
    doc = Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs])
    
    # 5. NO BLOCKING CONDITION
    assert "NO EXTENSION SHALL BE GRANTED UNTIL ALL BLOCKING CONDITIONS ARE RESOLVED." not in text
    assert "No blocking conditions identified." in text
    
    # 7. MISSING DATA
    assert "Not provided in source material." in text
    assert "No agent checks provided" in text
    assert "No findings provided" in text
    
    print("TEST 5, 7 PASSED: Missing Data and No Blocking Conditions")

def test_tool_registry():
    from app.tools.registry import tool_registry
    tool = tool_registry.get_tool("generate_approval_document")
    assert tool is not None
    print("TEST 9 PASSED: Tool Registry")

def test_download_api():
    from app.services.document_generator import document_generator_service
    task_id = "test_task_789_v2"
    document_generator_service.generate_approval_document(task_id, MOCK_DOC_DATA)
    
    res = client.get(f"/api/tasks/{task_id}/download")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert len(res.content) > 0
    print("TEST 11 PASSED: Download API")

def test_path_traversal():
    import asyncio
    from app.api.routes.tasks import download_task_output
    
    # Test path traversal directly against the endpoint function
    res = asyncio.run(download_task_output("../../windows/system32"))
    assert "error" in res
    assert res["error"]["code"] == "OUTPUT_INVALID"
    print("TEST 13 PASSED: Path Traversal")

def test_task_integration():
    # 10. TASK INTEGRATION & M12.1 E2E Flow
    # First upload a dummy document so RAG has something to find
    doc_content = """Inspection Report
Equipment: Tank T-204
Minimum allowable thickness: 10.0 mm
Measurements:
P1: 12.4 mm
P2: 11.8 mm
P3: 10.9 mm
P4: 12.1 mm
P5: 11.5 mm
Findings: Minor corrosion observed near P3.
Status: Review required."""
    
    files = {"file": ("test_permit.txt", doc_content.encode('utf-8'), "text/plain")}
    up_res = client.post("/api/documents/upload", files=files)
    doc_id = up_res.json()["document_id"]
    
    res = client.post("/api/tasks", json={
        "user_request": "Analyze this document and prepare an approval note based only on the information in the uploaded document.",
        "document_id": doc_id
    })
    task_id = res.json()["task_id"]
    
    # Run task
    run_res = client.post(f"/api/tasks/{task_id}/run", timeout=200.0)
    assert run_res.status_code == 200
    
    data = run_res.json()
    if data["status"] != "completed":
        log_res = client.get(f"/api/tasks/{task_id}/agent-log")
        print("Agent log:", log_res.json())
        print("Integration failed with data:", data)
    assert data["status"] == "completed"
    
    # Verify result contains output metadata (M12.1)
    result = data.get("result", {})
    assert "output" in result, "Task result missing output metadata"
    assert result["output"]["type"] == "docx"
    assert "path" in result["output"]
    
    # Download test
    dl_res = client.get(f"/api/tasks/{task_id}/download")
    assert dl_res.status_code == 200
    assert dl_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    
    # Validate actual downloaded content
    file_path = f"test_download_{task_id}.docx"
    with open(file_path, "wb") as f:
        f.write(dl_res.content)
        
    validate_docx(file_path)
    print(f"Generated E2E DOCX at {file_path}")
    # os.remove(file_path)
    
    print("TEST 10 PASSED: Task Integration & E2E Output")

if __name__ == "__main__":
    test_generate_and_validate_document()
    test_missing_data()
    test_tool_registry()
    test_download_api()
    test_path_traversal()
    test_task_integration()
    print("ALL M12 TESTS PASSED")
