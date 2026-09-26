import os
import uuid
import logging
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.config import settings

logger = logging.getLogger(__name__)

class DocumentGeneratorService:
    def __init__(self):
        self.outputs_dir = os.path.join(settings.storage_root, "outputs")
        os.makedirs(self.outputs_dir, exist_ok=True)

    def _set_status_style(self, run, status: str):
        status = status.upper()
        run.bold = True
        if status in ["PASS", "CONFIRMED"]:
            run.font.color.rgb = RGBColor(0, 128, 0) # Green
        elif status in ["FAILED", "FAIL", "EXPIRED", "INVALID", "BLOCKED"]:
            run.font.color.rgb = RGBColor(255, 0, 0) # Red
        elif status in ["WARNING", "PENDING"]:
            run.font.color.rgb = RGBColor(255, 165, 0) # Orange
        else:
            run.font.color.rgb = RGBColor(128, 128, 128) # Grey

    def generate_approval_document(self, task_id: str, data: dict) -> str:
        """
        Generates an industrial approval DOCX based on structured data.
        Returns the absolute file path of the generated document.
        """
        task_output_dir = os.path.join(self.outputs_dir, task_id)
        # Security: prevent path traversal
        if not os.path.abspath(task_output_dir).startswith(os.path.abspath(self.outputs_dir)):
            raise ValueError("Invalid task ID path traversal.")
        
        os.makedirs(task_output_dir, exist_ok=True)
        filename = "Industrial_Authorization.docx"
        file_path = os.path.join(task_output_dir, filename)

        doc = Document()
        
        # Safe fallback for get
        def safe_get(d, key, default="Not provided in source material."):
            return d.get(key, default) if d.get(key) else default

        # --- PAGE 1: HEADER & METADATA ---
        org = data.get("organization")
        if not org or org == "Not provided in source material." or "alibaba" in org.lower():
            org = "Builder AI Workbench"
        doc.add_heading(org, 0)
        doc.add_heading(safe_get(data, "title", "INDUSTRIAL AUTHORIZATION DOCUMENT"), 1)

        doc.add_paragraph(f"Document ID: {safe_get(data, 'document_id')}")
        if data.get("extension_id"):
            doc.add_paragraph(f"Extension ID: {data['extension_id']}")
        doc.add_paragraph(f"Permit Reference: {safe_get(data, 'permit_reference')}")
        doc.add_paragraph(f"Equipment: {safe_get(data, 'equipment')}")
        if "source_report" in data:
            doc.add_paragraph(f"Source Inspection Report: {safe_get(data, 'source_report')}")
            doc.add_paragraph(f"Inspection Date: {safe_get(data, 'inspection_date')}")
            doc.add_paragraph(f"Inspection Method: {safe_get(data, 'inspection_method')}")
        doc.add_paragraph(f"Classification: {safe_get(data, 'classification', 'CONFIDENTIAL')}")
        doc.add_paragraph(f"Timestamp: {safe_get(data, 'timestamp', datetime.utcnow().isoformat() + 'Z')}")

        doc.add_heading("AI AGENT SAFETY VERIFICATION & DRAFT SUMMARY", level=2)
        
        status_p = doc.add_paragraph("Status: ")
        status_run = status_p.add_run(f"DRAFT — AI-GENERATED — {safe_get(data, 'status', 'PENDING')} SIGN-OFF")
        status_run.bold = True
        
        warning_p = doc.add_paragraph()
        warning_run = warning_p.add_run("AI-DRAFTED — Requires human review and sign-off by the authorized Approver role before this document becomes valid. This document has no legal/operational authorization until approved and signed by the designated human authority.")
        warning_run.bold = True
        warning_run.font.color.rgb = RGBColor(255, 0, 0)
        
        doc.add_heading("PERMIT DETAILS", level=2)
        permit_details = data.get("permit_details", {})
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        
        details = [
            ("Work Area", safe_get(permit_details, "work_area")),
            ("Requested Extension", safe_get(permit_details, "requested_extension")),
            ("Current Status", safe_get(permit_details, "current_status")),
            ("Requested By", safe_get(permit_details, "requested_by")),
            ("Approver Role", safe_get(permit_details, "approver_role"))
        ]
        
        for key, val in details:
            row_cells = table.add_row().cells
            row_cells[0].text = key
            row_cells[1].text = val

        doc.add_page_break()

        # --- PAGE 2: AGENT CHECKS & FINDINGS ---
        doc.add_heading("AGENT CHECK RESULTS SUMMARY", level=2)
        agent_checks = data.get("agent_checks", [])
        if not agent_checks:
            doc.add_paragraph("No agent checks provided in source material.")
        else:
            for i, check in enumerate(agent_checks):
                doc.add_heading(f"{i+1}. {safe_get(check, 'name', 'Unnamed Check').upper()}", level=3)
                doc.add_paragraph(f"Check ID: {safe_get(check, 'check_id')}")
                
                res_p = doc.add_paragraph("Result: ")
                status = safe_get(check, 'status', 'NOT_VERIFIED')
                res_run = res_p.add_run(status)
                self._set_status_style(res_run, status)
                
                doc.add_paragraph(f"Evidence: {safe_get(check, 'evidence')}")
                doc.add_paragraph(f"Reference: {safe_get(check, 'reference')}")

        for table_data in data.get("measurements", []):
            doc.add_heading("INSPECTION MEASUREMENTS (VERBATIM FROM SOURCE)", level=2)
            src = table_data.get("source", {})
            doc.add_paragraph(f"Source: {src.get('document')} (page {src.get('page') or 'N/A'}, {src.get('chunk_id')})")
            columns = ["Point"] + list(table_data.get("columns", []))
            m_table = doc.add_table(rows=1, cols=len(columns))
            m_table.style = 'Table Grid'
            for i, col in enumerate(columns):
                m_table.rows[0].cells[i].text = col
            for row in table_data.get("rows", []):
                cells = m_table.add_row().cells
                for i, val in enumerate([row.get("point", "")] + list(row.get("values", []))):
                    if i < len(cells):
                        cells[i].text = str(val)

        doc.add_heading("FINDINGS SUMMARY", level=2)
        findings = data.get("findings", [])
        if not findings:
            doc.add_paragraph("No findings provided in source material.")
        else:
            for finding in findings:
                doc.add_paragraph(f"Finding: {safe_get(finding, 'finding_id', 'Unknown')}")
                doc.add_paragraph(f"Observation: {safe_get(finding, 'observation')}")
                doc.add_paragraph(f"Evidence: {safe_get(finding, 'evidence')}")
                doc.add_paragraph(f"Reference: {safe_get(finding, 'reference')}")
                
        doc.add_page_break()

        # --- PAGE 3: REFERENCES & BLOCKING CONDITIONS ---
        doc.add_heading("REGULATORY REFERENCE SUMMARY", level=2)
        references = data.get("regulatory_references", [])
        if not references:
            doc.add_paragraph("Reference not available in supplied source material.")
        else:
            for i, ref in enumerate(references):
                doc.add_paragraph(f"[{i+1}] Reference: {safe_get(ref, 'document')}")
                doc.add_paragraph(f"Section: {safe_get(ref, 'section')} | Page: {safe_get(ref, 'page', 'N/A')}")
                doc.add_paragraph(f"Requirement: {safe_get(ref, 'requirement')}")
                doc.add_paragraph(f"Application: {safe_get(ref, 'application')}")

        doc.add_heading("BLOCKING CONDITIONS", level=2)
        blocking = data.get("blocking_conditions", [])
        if not blocking:
            doc.add_paragraph("No blocking conditions identified.")
        else:
            for i, block in enumerate(blocking):
                doc.add_heading(f"BLOCKING CONDITION #{i+1}", level=3)
                doc.add_paragraph(f"Condition: {safe_get(block, 'condition')}")
                
                bp = doc.add_paragraph("Status: ")
                br = bp.add_run("BLOCKING")
                br.bold = True
                br.font.color.rgb = RGBColor(255, 0, 0)
                
                doc.add_paragraph(f"Reason: {safe_get(block, 'reason')}")
                doc.add_paragraph(f"Evidence: {safe_get(block, 'evidence')}")
                doc.add_paragraph(f"Required Action: {safe_get(block, 'required_action')}")
                doc.add_paragraph(f"Reference: {safe_get(block, 'reference')}")
            
            p = doc.add_paragraph()
            r = p.add_run("NO EXTENSION SHALL BE GRANTED UNTIL ALL BLOCKING CONDITIONS ARE RESOLVED.")
            r.bold = True
            r.font.color.rgb = RGBColor(255, 0, 0)

        doc.add_page_break()

        # --- PAGE 4: ACTIONS & APPROVAL ---
        doc.add_heading("ACTION REQUIRED", level=2)
        actions = data.get("required_actions", [])
        if not actions:
            doc.add_paragraph("No specific actions required in supplied source material.")
        else:
            for i, act in enumerate(actions):
                doc.add_paragraph(f"{i+1}. {safe_get(act, 'description')} (Role: {safe_get(act, 'responsible_role')})")

        doc.add_heading("AUTHORITY AND APPROVAL", level=2)
        doc.add_heading("Approval Conditions", level=3)
        doc.add_paragraph("The following conditions must be satisfied before final authorization:")
        doc.add_paragraph("☐ All blocking conditions resolved")
        doc.add_paragraph("☐ Required inspections completed")
        doc.add_paragraph("☐ Required safety checks verified")
        doc.add_paragraph("☐ Supporting evidence attached")
        doc.add_paragraph("☐ Authorized person reviewed the request")

        doc.add_heading("APPROVER DETAILS", level=3)
        approval = data.get("approval", {})
        doc.add_paragraph(f"Approver Role: {safe_get(approval, 'approver_role')}")
        doc.add_paragraph("Approver Name: ____________________________")
        doc.add_paragraph("Employee ID: ____________________________")
        doc.add_paragraph("Signature: ____________________________")
        doc.add_paragraph("Date: ____________________________")
        doc.add_paragraph("Time: ____________________________")
        
        doc.add_paragraph()
        dec_p = doc.add_paragraph("Decision (NOT YET APPROVED):")
        dec_p.bold = True
        doc.add_paragraph("☐ APPROVED")
        doc.add_paragraph("☐ REJECTED")
        doc.add_paragraph("☐ RETURNED FOR CORRECTION")
        doc.add_paragraph("If rejected, reason: ____________________________________________")
        
        doc.add_paragraph()
        end_p = doc.add_paragraph()
        end_r = end_p.add_run("END OF AI-DRAFTED AUTHORIZATION REQUEST\nDOCUMENT IS NOT VALID WITHOUT REQUIRED HUMAN REVIEW AND SIGN-OFF.")
        end_r.bold = True
        end_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_page_break()

        # --- PAGE 5: PROVENANCE ---
        doc.add_heading("DOCUMENT PROVENANCE", level=2)
        prov = data.get("provenance", {})
        doc.add_paragraph(f"Generated By: {safe_get(prov, 'agent', 'Sovereign AI Workbench')}")
        doc.add_paragraph(f"Task ID: {task_id}")
        
        doc.add_paragraph("Source Documents:")
        for doc_id in prov.get("source_documents", []):
            doc.add_paragraph(f"- {doc_id}")
            
        for warning in prov.get("warnings", []):
            doc.add_paragraph(f"Generation warning ({warning.get('stage')}): {warning.get('code')} — {warning.get('message')}")
        rejected = prov.get("rejected_values", [])
        if rejected:
            doc.add_paragraph("Values rejected by grounding guardrail (not supported by source):")
            for r in rejected:
                doc.add_paragraph(f"- [{r.get('field')}] {r.get('value')}")
        doc.add_paragraph(f"Generation Timestamp: {datetime.utcnow().isoformat()}Z")
        doc.add_paragraph("AI Draft Status: PENDING HUMAN REVIEW")

        doc.save(file_path)
        logger.info(f"Generated document saved to {file_path}")
        return file_path

document_generator_service = DocumentGeneratorService()
