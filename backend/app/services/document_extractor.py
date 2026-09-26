import os
import json
import csv
import logging
from typing import Dict, Any

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

from app.services.document_service import document_service

logger = logging.getLogger(__name__)

class DocumentExtractor:
    def extract(self, document_id: str) -> Dict[str, Any]:
        """
        Extracts content from a document and caches it.
        Returns the normalized content dictionary.
        """
        meta_result = document_service.get_metadata(document_id)
        if "error" in meta_result:
            return meta_result

        doc_dir = os.path.dirname(meta_result["storage_path"])
        content_cache_path = os.path.join(doc_dir, "content.json")

        # 1. Check Cache
        if os.path.exists(content_cache_path):
            try:
                with open(content_cache_path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception as e:
                logger.warning(f"Failed to read cache for {document_id}: {e}")
                # Fallback to re-extract if cache read fails

        file_path = meta_result["storage_path"]
        ext = meta_result["extension"].lower()

        # Base result skeleton
        result = {
            "document_id": document_id,
            "document_type": ext.strip("."),
            "content_status": "extracted",
            "requires_ocr": False,
            "text": "",
            "metadata": {}
        }

        # 2. Extract based on format
        try:
            if ext == ".pdf":
                self._extract_pdf(file_path, result)
            elif ext == ".docx":
                self._extract_docx(file_path, result)
            elif ext == ".txt":
                self._extract_txt(file_path, result)
            elif ext == ".csv":
                self._extract_csv(file_path, result)
            elif ext == ".xlsx":
                self._extract_xlsx(file_path, result)
            else:
                return {"error": {"code": "UNSUPPORTED_DOCUMENT_TYPE", "message": f"Cannot extract content from {ext} files."}}
        except Exception as e:
            logger.error(f"Extraction failed for {document_id}: {e}")
            return {"error": {"code": "EXTRACTION_FAILED", "message": "Failed to extract content from document."}}

        # 3. Post-process empty checks
        if result["content_status"] == "extracted" and not result["text"].strip():
            result["content_status"] = "empty"

        # 4. Cache the result
        try:
            with open(content_cache_path, "w", encoding="utf-8") as f:
                json.dump(result, f, indent=2, ensure_ascii=False)
        except Exception as e:
            logger.error(f"Failed to write cache for {document_id}: {e}")

        # 5. Update metadata
        document_service.update_metadata(document_id, {
            "content_status": result["content_status"],
            "requires_ocr": result["requires_ocr"]
        })

        return result

    def _extract_pdf(self, file_path: str, result: Dict[str, Any]):
        if not fitz:
            raise RuntimeError("PyMuPDF is not installed")
            
        doc = fitz.open(file_path)
        result["page_count"] = len(doc)
        result["pages"] = []
        
        full_text = []
        for i in range(len(doc)):
            page = doc.load_page(i)
            text = page.get_text() or ""
            text = text.strip()
            
            result["pages"].append({
                "page_number": i + 1,
                "text": text
            })
            if text:
                full_text.append(text)
                
        doc.close()
        
        combined_text = "\n\n".join(full_text).strip()
        result["text"] = combined_text
        
        if not combined_text and result["page_count"] > 0:
            result["content_status"] = "no_extractable_text"
            result["requires_ocr"] = True

    def _extract_docx(self, file_path: str, result: Dict[str, Any]):
        if not docx:
            raise RuntimeError("python-docx is not installed")
            
        doc = docx.Document(file_path)
        result["paragraphs"] = []
        full_text = []
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                result["paragraphs"].append(text)
                full_text.append(text)
                
        result["tables"] = []
        for table in doc.tables:
            table_data = {"rows": []}
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data["rows"].append(row_data)
                
                # Also append to full text logically
                full_text.append(" | ".join(row_data))
                
            result["tables"].append(table_data)
            
        result["text"] = "\n\n".join(full_text).strip()

    def _extract_txt(self, file_path: str, result: Dict[str, Any]):
        encodings = ['utf-8', 'latin-1', 'cp1252']
        text = ""
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    text = f.read()
                break
            except UnicodeDecodeError:
                continue
                
        result["text"] = text.strip()

    def _extract_csv(self, file_path: str, result: Dict[str, Any]):
        result["headers"] = []
        result["rows"] = []
        full_text = []
        
        with open(file_path, "r", encoding="utf-8", newline='') as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if i == 0:
                    result["headers"] = row
                    full_text.append(" | ".join(row))
                else:
                    result["rows"].append(row)
                    full_text.append(" | ".join(row))
                    
                # Safe limit to avoid massive in-memory arrays for MVP
                if i >= 1000:
                    full_text.append("... [truncated after 1000 rows]")
                    break
                    
        result["text"] = "\n".join(full_text)

    def _extract_xlsx(self, file_path: str, result: Dict[str, Any]):
        if not openpyxl:
            raise RuntimeError("openpyxl is not installed")
            
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        result["sheets"] = []
        full_text = []
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_data = {"name": sheet_name, "rows": []}
            full_text.append(f"Sheet: {sheet_name}")
            
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                row_values = [str(cell).strip() if cell is not None else "" for cell in row]
                # skip completely empty rows
                if any(row_values):
                    sheet_data["rows"].append(row_values)
                    full_text.append(" | ".join(row_values))
                    
                if i >= 1000:
                    full_text.append("... [truncated after 1000 rows]")
                    break
                    
            result["sheets"].append(sheet_data)
            full_text.append("") # spacer
            
        wb.close()
        result["text"] = "\n".join(full_text).strip()

document_extractor = DocumentExtractor()
